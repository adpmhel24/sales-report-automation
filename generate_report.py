
import csv
from collections import defaultdict
from datetime import date
from decimal import Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

BASE_DIR = Path(__file__).resolve().parent
INPUT_FILE = BASE_DIR / "sales_data.csv"
OUTPUT_FILE = BASE_DIR / "sales_report.xlsx"

CURRENCY_FORMAT = '"₱"#,##0.00'


def load_sales():
    records = []
    seen_ids = set()

    with INPUT_FILE.open("r", newline="", encoding="utf-8-sig") as file:
        reader = csv.DictReader(file)

        required = {
            "sale_id", "date", "product", "quantity",
            "unit_price", "payment_method"
        }

        if not reader.fieldnames or not required.issubset(reader.fieldnames):
            raise ValueError("CSV is missing required columns.")

        for line_number, row in enumerate(reader, start=2):
            try:
                sale_id = row["sale_id"].strip()
                product = row["product"].strip()
                payment = row["payment_method"].strip()
                sale_date = date.fromisoformat(row["date"].strip())
                quantity = int(row["quantity"])
                unit_price = Decimal(row["unit_price"])

                if not sale_id or not product or not payment:
                    raise ValueError("Missing transaction information")

                if sale_id in seen_ids:
                    raise ValueError(f"Duplicate sale ID: {sale_id}")

                if quantity <= 0 or unit_price < 0:
                    raise ValueError("Invalid quantity or unit price")

                if not unit_price.is_finite():
                    raise ValueError("Unit price must be finite")

                seen_ids.add(sale_id)

                records.append({
                    "id": sale_id,
                    "date": sale_date,
                    "product": product,
                    "quantity": quantity,
                    "price": unit_price,
                    "payment": payment,
                    "total": quantity * unit_price,
                })

            except (ValueError, InvalidOperation, TypeError, AttributeError) as exc:
                raise ValueError(
                    f"Invalid CSV data on line {line_number}: {exc}"
                ) from exc

    if not records:
        raise ValueError("CSV contains no sales records.")

    return records


def create_sheet(wb, title, headers, rows, money_columns=()):
    ws = wb.create_sheet(title)
    ws.append(headers)

    for row in rows:
        ws.append(row)

    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="17365D")
        cell.alignment = Alignment(horizontal="center")

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    for column in ws.columns:
        col_index = column[0].column
        max_length = max(
            len(str(cell.value or "")) for cell in column
        )

        ws.column_dimensions[
            get_column_letter(col_index)
        ].width = min(max(max_length + 3, 14), 42)

    for row in ws.iter_rows(min_row=2):
        for index in money_columns:
            row[index - 1].number_format = CURRENCY_FORMAT

    return ws


def generate_report(sales):
    daily = defaultdict(lambda: Decimal("0"))
    products = defaultdict(
        lambda: {"quantity": 0, "revenue": Decimal("0")}
    )
    payments = defaultdict(lambda: Decimal("0"))

    for sale in sales:
        daily[sale["date"]] += sale["total"]

        products[sale["product"]]["quantity"] += sale["quantity"]
        products[sale["product"]]["revenue"] += sale["total"]

        payments[sale["payment"]] += sale["total"]

    total_revenue = sum(
        (sale["total"] for sale in sales), Decimal("0")
    )
    total_quantity = sum(sale["quantity"] for sale in sales)
    total_transactions = len(sales)
    average_sale = total_revenue / total_transactions

    wb = Workbook()
    wb.remove(wb.active)

    create_sheet(
        wb,
        "Summary",
        ["Metric", "Value"],
        [
            ["Gross Sales", float(total_revenue)],
            ["Transactions", total_transactions],
            ["Units Sold", total_quantity],
            ["Average Transaction", float(average_sale)],
        ],
    )

    ws = wb["Summary"]
    ws["B2"].number_format = CURRENCY_FORMAT
    ws["B5"].number_format = CURRENCY_FORMAT

    create_sheet(
        wb,
        "Daily Sales",
        ["Date", "Gross Sales"],
        [
            [day.isoformat(), float(amount)]
            for day, amount in sorted(daily.items())
        ],
        money_columns=(2,),
    )

    create_sheet(
        wb,
        "Product Sales",
        ["Product", "Quantity", "Gross Sales"],
        [
            [name, data["quantity"], float(data["revenue"])]
            for name, data in sorted(products.items())
        ],
        money_columns=(3,),
    )

    create_sheet(
        wb,
        "Payment Summary",
        ["Payment Method", "Gross Sales"],
        [
            [method, float(amount)]
            for method, amount in sorted(payments.items())
        ],
        money_columns=(2,),
    )

    wb.save(OUTPUT_FILE)

    print(f"Report created: {OUTPUT_FILE}")
    print(f"Transactions: {total_transactions}")
    print(f"Units sold: {total_quantity}")
    print(f"Gross sales: PHP {total_revenue:,.2f}")


if __name__ == "__main__":
    try:
        sales = load_sales()
        generate_report(sales)
    except (FileNotFoundError, PermissionError, ValueError) as exc:
        print(f"Error: {exc}")
        raise SystemExit(1)
